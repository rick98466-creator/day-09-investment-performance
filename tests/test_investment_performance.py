import pandas as pd
import pytest
from openpyxl import load_workbook

from investment_performance.config import REPORTS_DIR
from investment_performance.data_loader import load_investment_prices
from investment_performance.validation import validate_investment_prices
from investment_performance.returns import calculate_returns
from investment_performance.risk import calculate_risk_metrics
from investment_performance.comparison import compare_investments
from investment_performance.reporting import export_performance_report


def test_dataset_loads():
    df = load_investment_prices()
    assert len(df) == 12
    assert df["date"].is_monotonic_increasing


def test_duplicate_dates_rejected():
    df = load_investment_prices()
    df.loc[df.index[1], "date"] = df.loc[df.index[0], "date"]
    with pytest.raises(ValueError, match="duplicated"):
        validate_investment_prices(df)


def test_nonpositive_prices_rejected():
    df = load_investment_prices()
    df.loc[df.index[0], "alpha_fund"] = 0
    with pytest.raises(ValueError, match="positive"):
        validate_investment_prices(df)


def test_return_calculations():
    df = pd.DataFrame({
        "date": pd.to_datetime(["2026-01-05", "2026-01-06"]),
        "alpha_fund": [100.0, 110.0],
    })
    daily, cumulative = calculate_returns(df)
    assert pd.isna(daily.loc[0, "alpha_fund"])
    assert daily.loc[1, "alpha_fund"] == pytest.approx(0.10)
    assert cumulative.loc[1, "alpha_fund"] == pytest.approx(0.10)


def test_risk_metrics_are_finite():
    results = calculate_risk_metrics(load_investment_prices())
    assert len(results) == 3
    assert results["annualized_volatility"].map(
        lambda x: pd.notna(x) and abs(x) != float("inf")
    ).all()
    assert (results["maximum_drawdown"] <= 0).all()


def test_comparison_ranked_by_return():
    results = compare_investments(load_investment_prices())
    assert results["rank"].tolist() == [1, 2, 3]
    assert results["cumulative_return"].is_monotonic_decreasing


def test_excel_report_has_four_sheets():
    export_performance_report(load_investment_prices())
    path = REPORTS_DIR / "investment_performance_report.xlsx"
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames == [
            "Executive Summary",
            "Daily Returns",
            "Cumulative Returns",
            "Risk Metrics",
        ]
    finally:
        workbook.close()